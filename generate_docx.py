from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_docx():
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'TH Sarabun PSK'
    font.size = Pt(16)
    font.color.rgb = RGBColor(0x2C, 0x2F, 0x33)

    # 1. Cover
    if os.path.exists('images/cover_image_1784879612091.jpg'):
        doc.add_picture('images/cover_image_1784879612091.jpg', width=Inches(6.0))
        
    title = doc.add_heading('📖 คู่มือการใช้งาน Discord', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.name = 'TH Sarabun PSK'
    title.runs[0].font.size = Pt(24)
    title.runs[0].font.color.rgb = RGBColor(0x58, 0x65, 0xF2)

    subtitle = doc.add_heading('🧪 Innovation Lab: ศน. โคราช', 1)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.name = 'TH Sarabun PSK'
    subtitle.runs[0].font.size = Pt(20)
    subtitle.runs[0].font.color.rgb = RGBColor(0x4F, 0x54, 0x5C)
    
    welcome_text = doc.add_paragraph('ยินดีต้อนรับท่านศึกษานิเทศก์ (ศน.) และกัลยาณมิตรทางวิชาการทุกท่านเข้าสู่ Innovation Lab: ศน. โคราช พื้นที่แห่งการเรียนรู้ แบ่งปัน และสร้างสรรค์นวัตกรรมการศึกษาของพวกเราค่ะ!')
    welcome_text.alignment = WD_ALIGN_PARAGRAPH.CENTER
    welcome_text.paragraph_format.space_after = Pt(14)

    # Admin Box Table
    admin_table = doc.add_table(rows=2, cols=1)
    admin_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_h = admin_table.rows[0].cells[0]
    cell_h.text = "🌟 คณะผู้ดูแลชุมชน (Community Administrators)"
    cell_h.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell_h.paragraphs[0].runs[0].font.name = 'TH Sarabun PSK'
    cell_h.paragraphs[0].runs[0].font.size = Pt(16)
    cell_h.paragraphs[0].runs[0].bold = True
    cell_h.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    set_cell_background(cell_h, '5865F2')

    cell_b = admin_table.rows[1].cells[0]
    p_adm = cell_b.paragraphs[0]
    p_adm.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_adm.paragraph_format.space_before = Pt(8)
    p_adm.paragraph_format.space_after = Pt(8)
    r1 = p_adm.add_run("👩‍🏫 ศน.ศิราณี กันชัย  |  👨‍🏫 ศน.แสนยากร สายสิน\n")
    r1.font.name = 'TH Sarabun PSK'
    r1.font.size = Pt(16)
    r1.bold = True
    r2 = p_adm.add_run("ศึกษานิเทศก์ชำนาญการพิเศษ\nสำนักงานศึกษาธิการจังหวัดนครราชสีมา")
    r2.font.name = 'TH Sarabun PSK'
    r2.font.size = Pt(15)
    r2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    set_cell_background(cell_b, 'F4F6FB')
    
    doc.add_page_break()

    # 2. Section 1
    if os.path.exists('images/channel_architecture_1784879622708.jpg'):
        doc.add_picture('images/channel_architecture_1784879622708.jpg', width=Inches(6.0))
    h1 = doc.add_heading('📂 1. แผนผังห้องต่างๆ (Channel Architecture)', level=1)
    h1.runs[0].font.name = 'TH Sarabun PSK'
    h1.runs[0].font.size = Pt(18)
    doc.add_paragraph('เซิร์ฟเวอร์ของเราถูกออกแบบมาเพื่อลดความซ้ำซ้อนและเน้นการใช้งานจริงค่ะ:')
    
    doc.add_heading('📌 หมวด Information (จุดประชาสัมพันธ์และคลังความรู้)', level=2)
    doc.add_paragraph('• #👋-ต้อนรับและกติกา: จุดเริ่มต้นสำหรับสมาชิกใหม่ (ท่านกำลังอ่านอยู่ที่นี่ค่ะ)')
    doc.add_paragraph('• #📢-ประกาศ-แจ้งข่าว: อัปเดตข่าวสารสำคัญ นโยบาย หรือประกาศจากส่วนกลางค่ะ')
    doc.add_paragraph('• #📚-คลังสื่อและนวัตกรรม: (รูปแบบฟอรั่ม) พื้นที่รวบรวมสื่อ แผนการสอน นวัตกรรม (มีระบบ Tag ค้นหาง่ายค่ะ)')
    doc.add_paragraph('• #🏆-โชว์ผลงาน: (รูปแบบแกลเลอรี) พื้นที่อวดโฉมผลงานที่ภาคภูมิใจ หรือความสำเร็จจากการลงพื้นที่นิเทศค่ะ')

    doc.add_heading('💬 หมวด Text Channels (ลานเสวนา)', level=2)
    doc.add_paragraph('• #☕-สภากาแฟ-คุยทั่วไป: ห้องนั่งเล่นสำหรับพูดคุยเรื่องทั่วไป แลกเปลี่ยนความคิดเห็นแบบสบายๆ ค่ะ')
    doc.add_paragraph('• #📅-นัดหมาย-ปรึกษางาน: พื้นที่สำหรับคุยเรื่องงานโดยเฉพาะ นัดลงพื้นที่ หรือหารือข้อราชการค่ะ')
    doc.add_paragraph('• #🎉-ห้องพักใจ-เล่าเรื่อง: พื้นที่ผ่อนคลาย คลายเครียดจากการทำงานค่ะ')
    doc.add_paragraph('• #ai-เพื่อนคู่คิด-ศึกษานิเทศก์: ห้องทำงานของ "ปราชญ์โคราช AI" ค่ะ')

    doc.add_heading('🎙️ หมวด Voice Channels (ห้องสนทนาเสียง)', level=2)
    doc.add_paragraph('• 🔊 ชวนกันเล่นเกม: ห้องสันทนาการหลังเลิกงานค่ะ')
    doc.add_paragraph('• 🔊 มุมกาแฟ (คุยเล่น): ห้องจิบกาแฟเปิดไมค์คุยแบบไม่เป็นทางการค่ะ')
    doc.add_paragraph('• 🔊 ห้องประชุม 1: สำหรับการประชุมงาน หรือแชร์หน้าจอพรีเซนต์งานค่ะ')
    
    doc.add_page_break()

    # 3. Section 2
    if os.path.exists('images/ai_tag_team_1784879632295.jpg'):
        doc.add_picture('images/ai_tag_team_1784879632295.jpg', width=Inches(6.0))
    h2 = doc.add_heading('🤖 2. ระบบผู้ช่วย AI แบบ Tag-Team', level=1)
    h2.runs[0].font.name = 'TH Sarabun PSK'
    h2.runs[0].font.size = Pt(18)
    doc.add_paragraph('ชุมชนของเรามี "ผู้ช่วย 2 รูปแบบ" ที่ทำงานร่วมกัน เพื่อยกระดับการทำงานวิชาการของท่านค่ะ:')
    
    doc.add_heading('💬 ผู้ช่วยคนที่ 1: "ปราชญ์โคราช AI"', level=2)
    doc.add_paragraph('(สถิตอยู่ในห้อง #ai-เพื่อนคู่คิด-ศึกษานิเทศก์ ค่ะ)')
    doc.add_paragraph('จุดเด่น: คิดไว ไอเดียพุ่ง เหมาะสำหรับ Brainstorm, ร่างโครงการ, หรือขอไอเดีย Active Learning ค่ะ')
    doc.add_paragraph('วิธีใช้: พิมพ์ @Korat Edu-Bot แล้วตามด้วยคำถาม หรือกดปุ่มเมนูลัดได้เลยค่ะ (ตัวอย่าง: @Korat Edu-Bot ช่วยออกแบบกิจกรรม Active Learning วิชาวิทยาศาสตร์ ป.4 ให้หน่อยค่ะ)')

    doc.add_heading('📚 ผู้ช่วยคนที่ 2: "ฐานข้อมูล NotebookLM"', level=2)
    doc.add_paragraph('(ลิงก์ปักหมุดอยู่ในห้อง #📚-คลังสื่อและนวัตกรรม ค่ะ)')
    doc.add_paragraph('จุดเด่น: แม่นยำ 100% ตอบจากอ้างอิงระเบียบราชการ เหมาะสำหรับค้นหาเกณฑ์ สมศ., ระเบียบ ว.PA, หรือคู่มือการทำงานต่างๆ ค่ะ')
    doc.add_paragraph('วิธีใช้: กดเข้าลิงก์ NotebookLM ที่ปักหมุดไว้ เพื่อสืบค้นข้อมูลเชิงลึกจากคลังเอกสารของจังหวัดได้ทันทีค่ะ')
    
    doc.add_page_break()

    # 4. Section 3 & 4
    if os.path.exists('images/community_culture_1784879641603.jpg'):
        doc.add_picture('images/community_culture_1784879641603.jpg', width=Inches(6.0))
    h3 = doc.add_heading('🤝 3. วัฒนธรรมของ Innovation Lab โคราช', level=1)
    h3.runs[0].font.name = 'TH Sarabun PSK'
    h3.runs[0].font.size = Pt(18)
    doc.add_paragraph('เราอยู่ร่วมกันด้วยวัฒนธรรม ไม่ใช่กฎระเบียบค่ะ:')
    doc.add_paragraph('1. พื้นที่ปลอดภัย (Safe Space): ไม่มีไอเดียไหนผิด กล้าที่จะแชร์ผลงานที่ "ยังไม่สมบูรณ์" เพื่อให้เพื่อนๆ ช่วยเติมเต็มค่ะ')
    doc.add_paragraph('2. แชร์แบบมีบริบท (Contextual Sharing): เวลาแปะลิงก์ผลงาน รบกวนเขียนอธิบายสั้นๆ 3 บรรทัด (ปัญหาคืออะไร? แก้อย่างไร? ผลลัพธ์เป็นอย่างไร?) เพื่อเป็นแรงบันดาลใจให้เพื่อนๆ ค่ะ')
    doc.add_paragraph('3. ทลายกำแพงสังกัด (Cross-Silo): ปัญหาของโรงเรียนเอกชน อาจแก้ได้ด้วยนวัตกรรมของ สพป. ขอให้เราช่วยเหลือกันในฐานะ "เพื่อนร่วมวิชาชีพ" โดยไม่แบ่งแยกสังกัดค่ะ')

    h4 = doc.add_heading('🎯 4. ภารกิจแรกของคุณ (First Quest!)', level=1)
    h4.runs[0].font.name = 'TH Sarabun PSK'
    h4.runs[0].font.size = Pt(18)
    doc.add_paragraph('เมื่อท่านอ่านคู่มือนี้จบแล้ว ขอเชิญแวะไปที่ห้อง #☕-สภากาแฟ-คุยทั่วไป นะคะ')
    doc.add_paragraph('พิมพ์แนะนำตัวสั้นๆ (ชื่อเล่น + สังกัด) เพื่อให้แอดมินทำการ "ติดยศและมอบสีประจำสังกัด" ให้กับท่านค่ะ:')
    
    doc.add_paragraph('- 🟦 สพป. (สีฟ้า)')
    doc.add_paragraph('- 🟩 สพม. (สีเขียว)')
    doc.add_paragraph('- 🟨 สช. (สีเหลือง)')
    doc.add_paragraph('- 🟧 อปท. (สีส้ม)')
    
    q = doc.add_paragraph('\n"ต่างบทบาท ต่างหน้าที่ แต่มีหัวใจดวงเดียวกัน...\nคือการยกระดับคุณภาพการศึกษาเพื่อเด็กไทยและลูกหลานชาวโคราชบ้านเราค่ะ" 🌾✨')
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.runs[0].font.name = 'TH Sarabun PSK'
    q.runs[0].font.size = Pt(17)
    q.runs[0].font.color.rgb = RGBColor(0xD3, 0x54, 0x00)
    q.runs[0].italic = True

    doc.save('Innovation_Lab_Manual.docx')
    print("Updated Innovation_Lab_Manual.docx successfully!")

if __name__ == '__main__':
    create_docx()
